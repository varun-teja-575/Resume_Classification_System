from docx import Document


def extract_text_from_docx(file_path):
    """
    Extract text from a DOCX resume.

    Parameters:
        file_path: Path to the DOCX file

    Returns:
        Extracted resume text
    """

    text = ""

    try:

        document = Document(
            file_path
        )

        # Extract paragraphs
        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                text += paragraph.text + "\n"


        # Extract tables
        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    if cell.text.strip():

                        row_text.append(
                            cell.text.strip()
                        )

                if row_text:

                    text += " ".join(
                        row_text
                    ) + "\n"


        return text.strip()

    except Exception as e:

        print("Error while reading DOCX:")
        print(e)

        return ""


if __name__ == "__main__":

    print("=" * 70)
    print("DOCX TEXT EXTRACTION TEST")
    print("=" * 70)

    file_path = input(
        "\nEnter DOCX file path: "
    )

    extracted_text = extract_text_from_docx(
        file_path
    )

    if extracted_text:

        print(
            "\nDOCX text extracted successfully!"
        )

        print("\n" + "=" * 70)
        print("EXTRACTED TEXT")
        print("=" * 70)

        print(
            extracted_text[:5000]
        )

    else:

        print(
            "\nNo text could be extracted from the DOCX."
        )